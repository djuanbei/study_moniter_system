"""DOCX export tests (PRD §79)."""

from __future__ import annotations

import io

from app.services.docx_export import question_set_docx


def _read_docx(blob: bytes):
    from docx import Document

    return Document(io.BytesIO(blob))


def test_docx_variants(db_session, sample_setup):
    qs = sample_setup["question_set"]
    for variant in ("student", "answer", "rubric"):
        blob = question_set_docx(db_session, qs, variant)
        assert blob[:2] == b"PK"  # OOXML zip magic
        doc = _read_docx(blob)
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "练习卷" in text
        assert "2x + 3 = 9" in text  # question prompt present
        if variant == "student":
            assert "答案：" not in text  # §78: student version has no answers
        else:
            assert "答案：" in text
        if variant == "rubric":
            assert "评分标准：" in text


def test_docx_score_table(db_session, sample_setup):
    doc = _read_docx(question_set_docx(db_session, sample_setup["question_set"], "student"))
    tables = doc.tables
    assert tables, "评分表 table should exist (§79 表格)"
    header = [c.text for c in tables[0].rows[0].cells]
    assert header[0] == "题号"
    assert "1" in header[1:]


def test_docx_rejects_unknown_variant(db_session, sample_setup):
    import pytest as _pytest

    with _pytest.raises(ValueError):
        question_set_docx(db_session, sample_setup["question_set"], "bogus")
